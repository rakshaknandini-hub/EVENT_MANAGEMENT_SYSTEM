import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

# Excel file name
FILE_NAME = "student_results.xlsx"

# Create Excel file if it does not exist
def create_excel():
    if not os.path.exists(FILE_NAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Results"

        headers = [
            "Name", "Roll No.", "Class",
            "Subject 1", "Subject 2", "Subject 3",
            "Subject 4", "Subject 5",
            "Total Marks", "Percentage", "Result"
        ]

        sheet.append(headers)
        workbook.save(FILE_NAME)


# Save student data
def save_student():
    try:
        name = name_entry.get()
        roll = roll_entry.get()
        student_class = class_entry.get()

        sub1 = int(sub1_entry.get())
        sub2 = int(sub2_entry.get())
        sub3 = int(sub3_entry.get())
        sub4 = int(sub4_entry.get())
        sub5 = int(sub5_entry.get())

        # Check empty fields
        if not name or not roll or not student_class:
            messagebox.showerror("Error", "Please fill all details.")
            return

        # Check marks
        marks = [sub1, sub2, sub3, sub4, sub5]

        if any(mark < 0 or mark > 100 for mark in marks):
            messagebox.showerror(
                "Error",
                "Marks should be between 0 and 100."
            )
            return

        # Calculate total and percentage
        total = sum(marks)
        percentage = total / 5

        # Pass if every subject has at least 40 marks
        if all(mark >= 40 for mark in marks):
            result = "Pass"
        else:
            result = "Fail"

        # Open Excel file
        workbook = load_workbook(FILE_NAME)
        sheet = workbook["Results"]

        # Check duplicate roll number
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if str(row[1]) == str(roll):
                messagebox.showerror(
                    "Error",
                    "Roll No. already exists."
                )
                return

        # Add data to Excel
        sheet.append([
            name,
            roll,
            student_class,
            sub1,
            sub2,
            sub3,
            sub4,
            sub5,
            total,
            f"{percentage:.1f}%",
            result
        ])

        workbook.save(FILE_NAME)

        messagebox.showinfo(
            "Success",
            "Student record saved successfully!"
        )

        clear_fields()

    except ValueError:
        messagebox.showerror(
            "Error",
            "Please enter valid numeric marks and Roll No."
        )


# Clear input fields
def clear_fields():
    entries = [
        name_entry, roll_entry, class_entry,
        sub1_entry, sub2_entry, sub3_entry,
        sub4_entry, sub5_entry
    ]

    for entry in entries:
        entry.delete(0, tk.END)


# Get result using Roll No.
def get_result():
    roll = result_roll_entry.get()

    if not roll:
        messagebox.showerror(
            "Error",
            "Please enter Roll No."
        )
        return

    workbook = load_workbook(FILE_NAME)
    sheet = workbook["Results"]

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[1]) == str(roll):

            found = True

            result_text = (
                f"Name: {row[0]}\n"
                f"Roll No.: {row[1]}\n"
                f"Class: {row[2]}\n"
                f"Total Marks: {row[8]}\n"
                f"Percentage: {row[9]}\n"
                f"Result: {row[10]}"
            )

            messagebox.showinfo(
                "Student Result",
                result_text
            )
            break

    if not found:
        messagebox.showerror(
            "Not Found",
            "Student record not found."
        )


# Show all results
def show_all_results():

    # Remove old data from Treeview
    for item in tree.get_children():
        tree.delete(item)

    workbook = load_workbook(FILE_NAME)
    sheet = workbook["Results"]

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if row[0] is not None:

            tree.insert(
                "",
                tk.END,
                values=(
                    row[0],
                    row[1],
                    row[2],
                    row[8],
                    row[9],
                    row[10]
                )
            )


# Exit application
def exit_program():
    root.destroy()


# ------------------------------------------------
# MAIN WINDOW
# ------------------------------------------------

create_excel()

root = tk.Tk()
root.title("Student Result Management System")
root.geometry("1000x650")

# Heading
title_label = tk.Label(
    root,
    text="STUDENT RESULT MANAGEMENT SYSTEM",
    font=("Arial", 20, "bold")
)

title_label.pack(pady=15)


# ------------------------------------------------
# ADD STUDENT FRAME
# ------------------------------------------------

add_frame = tk.LabelFrame(
    root,
    text="Add Student",
    font=("Arial", 12, "bold"),
    padx=15,
    pady=15
)

add_frame.pack(fill="x", padx=20, pady=10)


# Labels and entries
tk.Label(add_frame, text="Name").grid(
    row=0, column=0, padx=10, pady=5
)

name_entry = tk.Entry(add_frame)
name_entry.grid(row=0, column=1, padx=10)


tk.Label(add_frame, text="Roll No.").grid(
    row=0, column=2, padx=10
)

roll_entry = tk.Entry(add_frame)
roll_entry.grid(row=0, column=3, padx=10)


tk.Label(add_frame, text="Class").grid(
    row=0, column=4, padx=10
)

class_entry = tk.Entry(add_frame)
class_entry.grid(row=0, column=5, padx=10)


# Subjects
tk.Label(add_frame, text="Subject 1").grid(
    row=1, column=0, padx=10, pady=5
)

sub1_entry = tk.Entry(add_frame)
sub1_entry.grid(row=1, column=1)


tk.Label(add_frame, text="Subject 2").grid(
    row=1, column=2, padx=10
)

sub2_entry = tk.Entry(add_frame)
sub2_entry.grid(row=1, column=3)


tk.Label(add_frame, text="Subject 3").grid(
    row=1, column=4, padx=10
)

sub3_entry = tk.Entry(add_frame)
sub3_entry.grid(row=1, column=5)


tk.Label(add_frame, text="Subject 4").grid(
    row=2, column=0, padx=10, pady=5
)

sub4_entry = tk.Entry(add_frame)
sub4_entry.grid(row=2, column=1)


tk.Label(add_frame, text="Subject 5").grid(
    row=2, column=2, padx=10
)

sub5_entry = tk.Entry(add_frame)
sub5_entry.grid(row=2, column=3)


# Save button
save_button = tk.Button(
    add_frame,
    text="💾 Save",
    command=save_student,
    width=15
)

save_button.grid(
    row=2,
    column=5,
    padx=10,
    pady=5
)


# ------------------------------------------------
# GET RESULT FRAME
# ------------------------------------------------

result_frame = tk.LabelFrame(
    root,
    text="Get Result",
    font=("Arial", 12, "bold"),
    padx=15,
    pady=10
)

result_frame.pack(fill="x", padx=20, pady=10)


tk.Label(
    result_frame,
    text="Enter Roll No."
).pack(side="left", padx=10)


result_roll_entry = tk.Entry(result_frame)
result_roll_entry.pack(side="left", padx=10)


get_button = tk.Button(
    result_frame,
    text="🔍 Get Result",
    command=get_result,
    width=15
)

get_button.pack(side="left", padx=10)


show_button = tk.Button(
    result_frame,
    text="📋 Show All Results",
    command=show_all_results,
    width=18
)

show_button.pack(side="left", padx=10)


exit_button = tk.Button(
    result_frame,
    text="Exit",
    command=exit_program,
    width=10
)

exit_button.pack(side="right", padx=10)


# ------------------------------------------------
# SHOW ALL RESULTS TABLE
# ------------------------------------------------

table_frame = tk.LabelFrame(
    root,
    text="All Student Results",
    font=("Arial", 12, "bold")
)

table_frame.pack(
    fill="both",
    expand=True,
    padx=20,
    pady=10
)


columns = (
    "Name",
    "Roll No.",
    "Class",
    "Total Marks",
    "Percentage",
    "Result"
)

tree = ttk.Treeview(
    table_frame,
    columns=columns,
    show="headings"
)

for column in columns:

    tree.heading(
        column,
        text=column
    )

    tree.column(
        column,
        width=130,
        anchor="center"
    )


tree.pack(
    fill="both",
    expand=True
)


# Start application
root.mainloop()



