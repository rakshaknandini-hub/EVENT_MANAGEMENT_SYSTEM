# Import the load_workbook function from the openpyxl library
# It is used to open and work with Excel files
from openpyxl import load_workbook


# Name of the Excel file where account information will be stored
file_name = "atm.xlsx"


# Try to open the existing Excel file
try:
    # Load the existing workbook
    wb = load_workbook(file_name)

    # Select the worksheet named "ATM"
    ws = wb["ATM"]

# If the Excel file does not exist, create a new workbook
except FileNotFoundError:

    # Create a new Excel workbook
    from openpyxl import Workbook
    wb = Workbook()

    # Select the active worksheet
    ws = wb.active

    # Rename the worksheet to "ATM"
    ws.title = "ATM"

    # Create column headings
    ws.append(["Name", "Address", "Phone No", "Balance"])

    # Save the newly created workbook
    wb.save(file_name)


# Function to display the main menu
def menu():

    # Display the title of the banking system
    print("\n===== BANK ACCOUNT MANAGEMENT SYSTEM =====")

    # Display option 1
    print("1. Create Account")

    # Display option 2
    print("2. Deposit Money")

    # Display option 3
    print("3. Withdraw Money")

    # Display option 4
    print("4. Show Balance")

    # Display option 5
    print("5. Exit")


# Function to create a new bank account
def new_account():

    # Ask the user to enter their name
    name = input("Enter Name: ")

    # Ask the user to enter their address
    address = input("Enter Address: ")

    # Ask the user to enter their phone number
    phone = input("Enter Phone No: ")

    # Ask the user to enter the initial balance
    # float() converts the entered value into a decimal number
    balance = float(input("Enter Initial Balance: "))

    # Add the new account information as a new row in Excel
    ws.append([name, address, phone, balance])

    # Save the updated workbook
    wb.save(file_name)

    # Display a success message
    print("Account created successfully!")


# Function to deposit money into an account
def deposit():

    # Ask the user to enter the amount to deposit
    amount = float(input("Enter Amount to Deposit: "))

    # Read the current balance from column 4 of row 2
    # Column 4 represents the Balance column
    current_balance = ws.cell(row=2, column=4).value

    # Add the deposited amount to the current balance
    new_balance = current_balance + amount

    # Update the balance in the Excel sheet
    ws.cell(row=2, column=4).value = new_balance

    # Save the updated balance
    wb.save(file_name)

    # Display the new balance
    print("Money deposited successfully!")
    print("New Balance:", new_balance)


# Function to withdraw money from an account
def withdraw():

    # Ask the user to enter the amount they want to withdraw
    amount = float(input("Enter Amount to Withdraw: "))

    # Read the current balance from Excel
    current_balance = ws.cell(row=2, column=4).value

    # Check whether the account has enough money
    if amount <= current_balance:

        # Subtract the withdrawal amount from the balance
        new_balance = current_balance - amount

        # Update the balance in Excel
        ws.cell(row=2, column=4).value = new_balance

        # Save the updated workbook
        wb.save(file_name)

        # Display success message
        print("Money withdrawn successfully!")

        # Display the remaining balance
        print("Remaining Balance:", new_balance)

    else:

        # Display this message if withdrawal amount
        # is greater than the available balance
        print("Insufficient Balance!")


# Function to display the account balance
def showbalance():

    # Read the balance from the Excel sheet
    balance = ws.cell(row=2, column=4).value

    # Display the current balance
    print("Current Balance:", balance)


# Infinite loop keeps displaying the menu
# until the user selects Exit
while True:

    # Call the menu function
    menu()

    # Ask the user to select an option
    choice = input("Enter your choice: ")

    # If the user selects 1
    if choice == "1":

        # Call the function to create a new account
        new_account()

    # If the user selects 2
    elif choice == "2":

        # Call the deposit function
        deposit()

    # If the user selects 3
    elif choice == "3":

        # Call the withdrawal function
        withdraw()

    # If the user selects 4
    elif choice == "4":

        # Call the function to display balance
        showbalance()

    # If the user selects 5
    elif choice == "5":

        # Display exit message
        print("Thank you for using Bank Account Management System!")

        # Stop the while loop
        break

    # If the user enters any other option
    else:

        # Display an error message
        print("Invalid choice. Please try again.")
